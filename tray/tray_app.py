"""VigilServe Tray — Windows system tray controller.

A lightweight tray application that monitors and controls the VigilServe
services (backend API and local agent). The backend also serves the built
frontend, so no separate frontend service is needed. Includes a tkinter control
panel for service status, individual control, and port configuration.
"""
import base64
import os
import sys
import io
import json
import re
import ssl
import time
import socket
import subprocess
import threading
import webbrowser
import winreg
from datetime import datetime
from pathlib import Path
from typing import TextIO

from PIL import Image, ImageDraw, ImageFont, ImageTk
import pystray
import psutil
import tkinter as tk
from tkinter import font as tkfont, ttk, messagebox

try:
    from docx import Document
except Exception:
    Document = None

from config import (
    SERVICES,
    APP_ICON,
    PYTHON_EXE,
    PYTHON_MISSING_HINT,
    get_ports,
    set_ports,
    build_services,
    get_server_ip,
    set_server_ip,
    list_local_ips,
)
from logger import get_logger, logger

try:
    from docs_bundle import DOCS_BUNDLE
except Exception:
    DOCS_BUNDLE = {}


TRAY_VERSION = "1.1.65"

SUPPORT_CARDS: list[dict] = [
    {
        "key": "wechat",
        "icon": "\U0001F49A",
        "title": "微信赞赏",
        "desc": "微信扫码赞赏，支持一下我吧～",
        "button": "留言 \U0001F4AC",
        "link": "https://vigilserve.com/contact.html",
        "accent": False,
    },
    {
        "key": "alipay",
        "icon": "\U0001F499",
        "title": "支付宝赞赏",
        "desc": "支付宝扫码赞赏，支持任意金额。",
        "button": "留言 \U0001F4AC",
        "link": "https://vigilserve.com/contact.html",
        "accent": False,
    },
    {
        "key": "afdian",
        "icon": "\u26A1",
        "title": "爱发电",
        "desc": "关注我的爱发电主页，可月度赞助、留言互动。",
        "button": "前往爱发电 \u2197",
        "link": "https://www.ifdian.net/a/VigilServe",
        "accent": True,
    },
]

SUPPORT_ACCENT = "#17c5d6"
SUPPORT_ACCENT_HOVER = "#12b2c2"
SUPPORT_BORDER = "#e6e8eb"


def _resource_dir() -> Path:
    """Return the directory where bundled resources live.

    PyInstaller onefile mode extracts datas into a temporary ``sys._MEIPASS``
    folder at runtime. In normal source execution the resources sit next to
    this module.
    """
    if getattr(sys, "frozen", False) and hasattr(sys, "_MEIPASS"):
        return Path(sys._MEIPASS)
    return Path(__file__).parent


def _docx_path(name: str) -> Path:
    """Return the path to a bundled .docx file (fallback for source runs)."""
    return _resource_dir() / "docs" / name


def _docx_filename(title: str) -> str:
    """Map a document tab title to its source .docx filename."""
    mapping = {
        "用户协议": "VigilServe用户协议.docx",
        "免责声明": "VigilServe免责声明.docx",
        "联系作者": "联系作者.docx",
        "支持作者": "支持作者.docx",
    }
    return mapping.get(title, f"{title}.docx")


def _load_support_content(title: str) -> tuple[list[dict], list[bytes]]:
    """Load text blocks and QR-code images for the support-author content.

    Prefer the embedded bundle so end users cannot alter the distributed
    content. Fall back to reading the external .docx during development.
    """
    images: list[bytes] = []
    if title in DOCS_BUNDLE:
        entry = DOCS_BUNDLE[title]
        for img_info in entry.get("images", []):
            try:
                images.append(base64.b64decode(img_info["data_base64"]))
            except Exception as exc:
                logger.warning(f"Failed to decode support image: {exc}")
        return entry.get("blocks", []), images

    docx_path = _docx_path(_docx_filename(title))
    if Document is None:
        return [{"type": "paragraph", "text": f"无法读取文档：缺少 python-docx 依赖\n{docx_path}"}], images
    try:
        doc = Document(docx_path)
        blocks = []
        for para in doc.paragraphs:
            if not para.text.strip() and not para.runs:
                continue
            blocks.append(_extract_paragraph_block(para))
        for table in doc.tables:
            blocks.append(_extract_table_block(table))
        from docx.oxml.ns import qn
        for para in doc.paragraphs:
            for run in para._p.iter():
                if run.tag.endswith("}drawing") or run.tag.endswith("drawing"):
                    for blip in run.iter():
                        if blip.tag.endswith("}blip") or blip.tag.endswith("blip"):
                            embed = blip.get(qn("r:embed"))
                            if embed:
                                try:
                                    images.append(doc.part.related_parts[embed].blob)
                                except Exception as exc:
                                    logger.warning(f"Failed to load support image: {exc}")
        return blocks, images
    except Exception as exc:
        return [{"type": "paragraph", "text": f"无法读取文档：{exc}\n{docx_path}"}], images


def _extract_paragraph_block(para) -> dict:
    """Convert a python-docx paragraph to a block dict."""
    runs = []
    for run in para.runs:
        if not run.text:
            continue
        runs.append(
            {
                "text": run.text,
                "bold": bool(run.bold),
                "italic": bool(run.italic),
                "underline": bool(run.underline),
            }
        )
    return {
        "type": "paragraph",
        "text": para.text,
        "runs": runs,
        "bold": bool(para.runs and all(r.bold for r in para.runs if r.text)),
        "align": _align_name(para.alignment),
        "first_line_indent": _pt(para.paragraph_format.first_line_indent),
        "left_indent": _pt(para.paragraph_format.left_indent),
    }


def _extract_table_block(table) -> dict:
    """Convert a python-docx table to a block dict."""
    rows = []
    for row in table.rows:
        rows.append([cell.text.strip() for cell in row.cells])
    return {"type": "table", "rows": rows}


def _pt(value) -> float | None:
    if value is None:
        return None
    try:
        return value.pt
    except Exception:
        return None


def _align_name(value) -> str | None:
    if value is None:
        return None
    mapping = {0: "left", 1: "center", 2: "right", 3: "justify"}
    return mapping.get(value, "left")


def _render_support_text(
    parent: tk.Frame,
    blocks: list[dict],
    wraplength: int = 760,
) -> None:
    """Render support-author paragraphs as stacked labels with simple formatting."""
    for block in blocks:
        if block.get("type") != "paragraph":
            continue
        text = block.get("text", "")
        if not text.strip():
            continue
        font = ("Microsoft YaHei", 11)
        if block.get("bold"):
            font = ("Microsoft YaHei", 11, "bold")
        align = block.get("align")
        anchor = "w"
        if align == "center":
            anchor = "center"
        elif align == "right":
            anchor = "e"
        tk.Label(
            parent,
            text=text,
            bg="#ffffff",
            fg="#333333",
            font=font,
            anchor=anchor,
            justify="left" if anchor != "center" else "center",
            wraplength=wraplength,
        ).pack(fill=tk.X, pady=(0, 6))


def _emoji_icon_photo(char: str, box: int = 38) -> "ImageTk.PhotoImage | None":
    """Render an emoji glyph into a transparent tkinter photo image.

    Windows ships the Segoe UI Emoji font with colour (COLR) glyphs; Pillow can
    rasterise them when ``embedded_color`` is enabled. Returns ``None`` when the
    font or the colour rasteriser is unavailable so the caller can fall back to
    a plain text label.
    """
    font_path = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / "seguiemj.ttf"
    try:
        font = ImageFont.truetype(str(font_path), box)
        canvas = Image.new("RGBA", (box * 3, box * 3), (0, 0, 0, 0))
        ImageDraw.Draw(canvas).text((box, box), char, font=font, embedded_color=True)
        bbox = canvas.getbbox()
        if bbox:
            canvas = canvas.crop(bbox)
        canvas.thumbnail((box, box), Image.LANCZOS)
        return ImageTk.PhotoImage(canvas)
    except Exception as exc:
        logger.warning(f"Emoji icon render failed ({char!r}): {exc}")
        return None


def _classify_qr(blob: bytes) -> "str | None":
    """Guess the payment platform of a QR image from its dominant brand colour."""
    try:
        img = Image.open(io.BytesIO(blob)).convert("RGB")
        width, height = img.size
        samples = [
            img.getpixel((int(width * 0.03), int(height * 0.02))),
            img.getpixel((int(width * 0.05), int(height * 0.03))),
            img.getpixel((int(width * 0.97), int(height * 0.02))),
        ]
        red = sum(p[0] for p in samples) / len(samples)
        green = sum(p[1] for p in samples) / len(samples)
        blue = sum(p[2] for p in samples) / len(samples)
        if green > red + 30 and green > blue + 30:
            return "wechat"
        if blue > red + 30 and blue > green + 30:
            return "alipay"
    except Exception as exc:
        logger.warning(f"QR image classification failed: {exc}")
    return None


def _support_qr_map(images: list[bytes]) -> dict[str, bytes]:
    """Map QR images to card keys, preferring colour detection over position."""
    mapping: dict[str, bytes] = {}
    for blob in images:
        kind = _classify_qr(blob)
        if kind and kind not in mapping:
            mapping[kind] = blob
    positional = ["alipay", "wechat"]
    for idx, blob in enumerate(images):
        if idx < len(positional) and positional[idx] not in mapping:
            mapping[positional[idx]] = blob
    return mapping


def _wrap_cjk(text: str, width: int, font) -> str:
    """Wrap text at a pixel width while keeping CJK punctuation off line starts.

    tkinter's ``wraplength`` breaks blindly, which can leave 。，、 sitting at the
    beginning of a line. This helper pulls the preceding character down with the
    punctuation so the result matches Chinese typographic conventions.
    """
    if not text:
        return text
    try:
        measure = tkfont.Font(font=font).measure
    except Exception:
        return text

    no_line_start = "，。、；：！？）】》」』”’%,.;:!?)]}"
    lines: list[str] = []
    current = ""
    current_width = 0

    for char in text:
        char_width = measure(char)
        if current and current_width + char_width > width:
            if char in no_line_start and len(current) > 1:
                lines.append(current[:-1])
                current = current[-1] + char
                current_width = measure(current)
                continue
            lines.append(current)
            current = char
            current_width = char_width
        else:
            current += char
            current_width += char_width
    if current:
        lines.append(current)
    return "\n".join(lines)


def _build_support_button(parent: tk.Widget, spec: dict) -> tk.Label:
    """Create a flat, clickable label that acts as the card's action button."""
    accent = bool(spec.get("accent"))
    normal_bg = SUPPORT_ACCENT if accent else "#ffffff"
    hover_bg = SUPPORT_ACCENT_HOVER if accent else "#f3f4f6"
    button = tk.Label(
        parent,
        text=spec["button"],
        font=("Microsoft YaHei", 10),
        bg=normal_bg,
        fg="#ffffff" if accent else "#24292f",
        bd=0,
        relief=tk.FLAT,
        padx=12,
        pady=7,
        cursor="hand2",
        highlightthickness=0 if accent else 1,
        highlightbackground="#d0d7de",
        highlightcolor="#d0d7de",
    )

    link = spec["link"]

    def _open_link(_event=None) -> None:
        try:
            logger.info(f"Open support link: {link}")
            webbrowser.open(link)
        except Exception as exc:
            logger.error(f"Open support link failed ({link}): {exc}")

    button.bind("<Button-1>", _open_link)
    button.bind("<Enter>", lambda _e: button.configure(bg=hover_bg))
    button.bind("<Leave>", lambda _e: button.configure(bg=normal_bg))
    return button


def _render_support_cards(
    parent: tk.Widget,
    images: list[bytes],
    image_refs: list,
) -> None:
    """Render the three donation cards, mirroring the official donation page."""
    qr_map = _support_qr_map(images)

    cards = tk.Frame(parent, bg="#ffffff")
    cards.pack(fill=tk.X, padx=18, pady=(2, 16))
    for col in range(len(SUPPORT_CARDS)):
        cards.grid_columnconfigure(col, weight=1, uniform="support-card")

    for col, spec in enumerate(SUPPORT_CARDS):
        border = tk.Frame(cards, bg=SUPPORT_BORDER)
        border.grid(row=0, column=col, sticky="nsew", padx=7, pady=2)
        inner = tk.Frame(border, bg="#ffffff")
        inner.pack(fill=tk.BOTH, expand=True, padx=1, pady=1)

        body = tk.Frame(inner, bg="#ffffff")
        body.pack(fill=tk.BOTH, expand=True, padx=18, pady=(14, 12))

        icon_photo = _emoji_icon_photo(spec["icon"], 34)
        if icon_photo is not None:
            image_refs.append(icon_photo)
            tk.Label(body, image=icon_photo, bg="#ffffff", bd=0).pack()
        else:
            tk.Label(
                body,
                text=spec["icon"],
                font=("Segoe UI Emoji", 18),
                bg="#ffffff",
                fg="#333333",
            ).pack()

        tk.Label(
            body,
            text=spec["title"],
            font=("Microsoft YaHei", 11, "bold"),
            bg="#ffffff",
            fg="#1f2328",
        ).pack(pady=(8, 6))

        desc_font = ("Microsoft YaHei", 9)
        tk.Label(
            body,
            text=_wrap_cjk(spec["desc"], 200, desc_font),
            font=desc_font,
            bg="#ffffff",
            fg="#6b7280",
            justify="center",
        ).pack()

        blob = qr_map.get(spec["key"])
        has_qr = False
        if blob:
            try:
                qr_img = Image.open(io.BytesIO(blob)).convert("RGB")
                qr_img.thumbnail((160, 168), Image.LANCZOS)
                qr_photo = ImageTk.PhotoImage(qr_img)
                image_refs.append(qr_photo)
                tk.Label(body, image=qr_photo, bg="#ffffff", bd=0).pack(pady=(12, 0))
                has_qr = True
            except Exception as exc:
                logger.warning(f"Failed to render QR card {spec['key']}: {exc}")

        if has_qr:
            _build_support_button(body, spec).pack(side=tk.BOTTOM, fill=tk.X, pady=(12, 0))
        else:
            tk.Frame(body, bg="#ffffff").pack(fill=tk.BOTH, expand=True)
            _build_support_button(body, spec).pack(fill=tk.X)
            tk.Frame(body, bg="#ffffff").pack(fill=tk.BOTH, expand=True)


def _render_support_content(
    parent: tk.Widget,
    blocks: list[dict],
    images: list[bytes],
    image_refs: list,
    wraplength: int = 760,
) -> None:
    """Render the full support-author panel (scrollable text + donation cards)."""
    canvas = tk.Canvas(parent, bg="#ffffff", highlightthickness=0)
    scrollbar = tk.Scrollbar(parent, orient=tk.VERTICAL, command=canvas.yview)
    canvas.configure(yscrollcommand=scrollbar.set)
    scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
    canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=10, pady=10)

    content = tk.Frame(canvas, bg="#ffffff")
    canvas_window = canvas.create_window((0, 0), window=content, anchor="nw")

    content.bind("<Configure>", lambda e: canvas.configure(scrollregion=canvas.bbox("all")))

    text_frame = tk.Frame(content, bg="#ffffff")
    text_frame.pack(fill=tk.X, padx=24, pady=(14, 6))
    _render_support_text(text_frame, blocks, wraplength=wraplength)

    _render_support_cards(content, images, image_refs)

    def _on_canvas_configure(event: tk.Event) -> None:
        canvas.itemconfig(canvas_window, width=event.width)

    canvas.bind("<Configure>", _on_canvas_configure)


def _service_log_file(service_id: str) -> Path:
    """Return the log file path for a service's stdout/stderr."""
    log_dir = Path(os.environ.get("LOCALAPPDATA", "")) / "VigilServe" / "Tray" / "logs" / datetime.now().strftime("%Y-%m-%d")
    log_dir.mkdir(parents=True, exist_ok=True)
    return log_dir / f"service.{service_id}.log"


def _read_log_tail(path: Path, limit: int = 1500) -> str:
    """Return the tail of a service log, for surfacing start-up failures.

    An empty file is itself diagnostic: it means the spawned process died
    before writing a single byte (e.g. a Microsoft Store `python.exe` alias).
    """
    try:
        text = Path(path).read_text(encoding="utf-8", errors="replace").strip()
    except Exception:
        return ""
    if not text:
        return "（日志为空：进程未产生任何输出）"
    return text[-limit:]


def _parse_requirements(path: Path) -> list:
    """Return ``[(distribution_name, specifier), ...]`` for a requirements file.

    Extras are stripped (``uvicorn[standard]`` -> ``uvicorn``) because
    ``importlib.metadata.version()`` only knows the distribution name.
    """
    reqs = []
    for raw in path.read_text(encoding="utf-8", errors="replace").splitlines():
        line = raw.split("#", 1)[0].strip()
        if not line or line.startswith("-"):
            continue
        m = re.match(r"^([A-Za-z0-9_.\-]+)\s*(?:\[[^\]]*\])?\s*(.*)$", line)
        if m:
            reqs.append((m.group(1), m.group(2).strip()))
    return reqs


# therefore absent from requirements.txt. Checking only the pinned names let a
_SMOKE_IMPORTS = (
    "fastapi",
    "uvicorn",
    "sqlalchemy",
    "apscheduler.schedulers.background",
    "websockets",
    "pydantic",
    "multipart",
    "jinja2",
    "winrm",
    "paramiko",
    "winpty",
    "psutil",
    "requests",
)


# the only honest test of "can this service start?" — metadata bookkeeping is not.
_SMOKE_SNIPPET = r"""
import importlib, json, sys

failed = []
for name in json.loads(sys.argv[1]):
    try:
        importlib.import_module(name)
    except BaseException as exc:
        failed.append([name, "%s: %s" % (type(exc).__name__, exc)])
try:
    from tzlocal import get_localzone
    get_localzone()          # needs the `tzdata` package on Windows
except BaseException as exc:
    failed.append(["tzlocal.get_localzone()", "%s: %s" % (type(exc).__name__, exc)])
print(json.dumps(failed, ensure_ascii=True))
"""


def _probe_imports(log):
    """Import the backend's dependency graph and return what fails.

    ``None`` means "could not determine" — callers then fall back to running
    pip, which is the safe direction.
    """
    if not PYTHON_EXE:
        return None
    try:
        proc = subprocess.run(
            [PYTHON_EXE, "-c", _SMOKE_SNIPPET, json.dumps(_SMOKE_IMPORTS)],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=300,
            creationflags=subprocess.CREATE_NO_WINDOW,
        )
    except Exception as exc:
        log.warning(f"Dependency import check could not run: {exc}")
        return None
    if proc.returncode != 0:
        log.warning(
            f"Dependency import check failed (code={proc.returncode}): "
            f"{(proc.stderr or '')[-400:]}"
        )
        return None

    lines = [l.strip() for l in (proc.stdout or "").splitlines() if l.strip()]
    if not lines:
        log.warning("Dependency import check produced no output")
        return None
    try:
        return json.loads(lines[-1])
    except Exception as exc:
        log.warning(f"Dependency import check output was not JSON: {exc}")
        return None


# Runs inside the *target* interpreter. Emits ASCII-only JSON so it survives
_PROBE_SNIPPET = r"""
import json, sys
from importlib.metadata import version, PackageNotFoundError

report = []
for name, spec in json.loads(sys.argv[1]):
    try:
        actual = version(name)
    except PackageNotFoundError:
        report.append([name, spec, None, False])
        continue
    ok = True
    if spec.startswith("=="):
        def norm(v):
            return tuple(int(x) for x in v.split("+")[0].split(".") if x.isdigit())
        ok = norm(actual) == norm(spec[2:].strip())
    report.append([name, spec, actual, ok])
print(json.dumps(report))
"""


def _probe_requirements(req: Path, log):
    """List the requirements that are missing or at the wrong version.

    ``None`` means "could not determine" — callers then fall back to running
    pip, which is the safe direction.
    """
    if not PYTHON_EXE:
        return None
    try:
        reqs = _parse_requirements(req)
    except Exception as exc:
        log.warning(f"Could not parse {req}: {exc}")
        return None
    if not reqs:
        return []

    try:
        proc = subprocess.run(
            [PYTHON_EXE, "-c", _PROBE_SNIPPET, json.dumps(reqs)],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=120,
            creationflags=subprocess.CREATE_NO_WINDOW,
        )
    except Exception as exc:
        log.warning(f"Dependency probe could not run: {exc}")
        return None
    if proc.returncode != 0:
        log.warning(
            f"Dependency probe failed (code={proc.returncode}): "
            f"{(proc.stderr or '')[-400:]}"
        )
        return None

    lines = [l.strip() for l in (proc.stdout or "").splitlines() if l.strip()]
    if not lines:
        log.warning("Dependency probe produced no output")
        return None
    try:
        report = json.loads(lines[-1])
    except Exception as exc:
        log.warning(f"Dependency probe output was not JSON: {exc}")
        return None

    return [
        {"name": name, "want": spec, "got": actual}
        for name, spec, actual, ok in report
        if not ok
    ]


def _touch_marker(marker: Path, log) -> None:
    """Memoise a successful dependency check. Never fatal if it fails."""
    try:
        marker.parent.mkdir(parents=True, exist_ok=True)
        marker.touch()
    except Exception as exc:
        log.warning(f"Could not write dependency marker {marker}: {exc}")


def _ensure_service_deps(service: dict, log) -> None:
    """Make sure a service's Python requirements are importable before start.

    Mirrors ``installer/scripts/start.bat``. The packaged backend is plain
    source and the tray is what actually launches it, so without this a fresh
    machine dies on the first start with ``ModuleNotFoundError: fastapi`` —
    while the tray happily reports a successful start.

    The installer now also ships a portable Python runtime with these packages
    already installed, so the expected path is: import the backend's dependency
    graph, find nothing broken, and never invoke pip. That matters because the
    bundled runtime lives under a read-only install path where ``pip install``
    cannot write. pip is only a fallback, e.g. when the backend gains a
    dependency after the runtime was built.

    The gate is an *import test*, not a version comparison, because the
    dependencies that break in practice are the transitive ones. A build that
    verified only the names in requirements.txt shipped a runtime without
    ``six`` and ``tzdata``; the tray announced "Dependencies already satisfied"
    and the backend exited with ``ModuleNotFoundError: six``.
    """
    deps = service.get("deps")
    if not deps:
        return

    marker = Path(deps["marker"])

    req = Path(deps["requirements"])
    if not req.is_file():
        log.warning(f"Dependency manifest missing, skipping install: {req}")
        return

    broken = _probe_imports(log)
    if broken is not None and not broken:
        log.info(f"Dependency check passed for {service['name']}")
        _touch_marker(marker, log)
        return

    if broken:
        for name, why in broken:
            log.error(f"Unusable dependency for {service['name']}: {name} -> {why}")
        missing = _probe_requirements(req, log)
        if missing:
            log.info("Version mismatches: " + ", ".join(
                f"{m['name']} (has {m['got'] or 'none'}, want {m['want'] or 'any'})"
                for m in missing
            ))
    else:
        log.info(f"Could not verify dependencies for {service['name']}; running pip")

    cmd = [PYTHON_EXE, "-m", "pip", "install", "--disable-pip-version-check", "-r", str(req)]
    manual = f'"{PYTHON_EXE}" -m pip install -r "{req}"'
    try:
        proc = subprocess.run(
            cmd,
            cwd=service.get("cwd") or os.getcwd(),
            capture_output=True,
            text=True,
            timeout=1800,
            creationflags=subprocess.CREATE_NO_WINDOW,
        )
    except Exception as exc:
        log.error(f"Dependency installation could not run: {exc}")
        raise RuntimeError(
            f"首次运行需要安装 {service['name']} 的 Python 依赖，但安装过程无法执行：\n{exc}\n\n"
            f"请手动执行后重试：\n  {manual}"
        ) from exc

    if proc.returncode != 0:
        tail = ((proc.stdout or "") + (proc.stderr or ""))[-1200:]
        log.error(f"Dependency installation failed (code={proc.returncode}): {tail}")
        raise RuntimeError(
            f"首次运行安装 {service['name']} 的依赖失败（代码 {proc.returncode}）。\n\n"
            f"请手动执行后重试：\n  {manual}\n\n"
            f"pip 输出：\n{tail}"
        )

    still_broken = _probe_imports(log)
    if still_broken:
        detail = "\n".join(f"  - {name}: {why}" for name, why in still_broken)
        raise RuntimeError(
            f"{service['name']} 的依赖已安装，但仍无法导入：\n{detail}\n\n"
            "若 VigilServe 安装在 Program Files 下，请以管理员身份运行托盘程序，"
            "或手动执行后重试：\n  " + manual
        )

    _touch_marker(marker, log)
    log.info(f"Dependencies installed for {service['name']}")



def _port_in_use(port: int) -> bool:
    """Fast check: try to connect to localhost:port."""
    try:
        with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
            s.settimeout(0.2)
            return s.connect_ex(("127.0.0.1", port)) == 0
    except Exception:
        return False


def _local_ca_path() -> str:
    """本机后端生成的本地 CA（只用于探明服务端是否已启用 HTTPS）。"""
    here = os.path.dirname(os.path.abspath(__file__))
    cands = [
        os.path.join(here, "..", "backend", "certs", "ca.crt"),
        os.path.join(here, "..", "..", "backend", "certs", "ca.crt"),
        os.path.join(here, "certs", "ca.crt"),
    ]
    for p in cands:
        p = os.path.normpath(p)
        if os.path.isfile(p):
            return p
    return ""


def _backend_scheme(port: int) -> str:
    """判断本机后端当前讲 http 还是 https（安全加固阶段 2 之后默认 https）。

    做法与 Agent 一致：试着做一次 TLS 握手，成功即 https。
    这里只用于拼浏览器地址，不涉及敏感数据，因此找不到 CA 时退化为
    “不校验证书只探握手”。

    ⚠ 交付审查 P1-6 复核（2026-09-23）：这里的 CERT_NONE **不是"关闭了通道加密"**，
    别照着"关闭了校验"的清单把它删掉。理由是上面那句 —— 本函数唯一职责是判断
    scheme；而且连接目标是 127.0.0.1 本机回环，对外部攻击者没有暴露面。
    有本地 CA 时它走的是 CERT_REQUIRED 真校验。
    """
    try:
        ctx = ssl.SSLContext(ssl.PROTOCOL_TLS_CLIENT)
        ctx.check_hostname = False
        ca = _local_ca_path()
        if ca:
            ctx.verify_mode = ssl.CERT_REQUIRED
            ctx.load_verify_locations(cafile=ca)
        else:
            # 只在"没有本地 CA"这一个分支退化：本函数要区分 http/https 就必须完成握手，
            # 严格校验会让"其实是 https"被误判成 http，结果地址拼错打不开。
            ctx.verify_mode = ssl.CERT_NONE
        with socket.create_connection(("127.0.0.1", int(port)), timeout=0.6) as raw:
            with ctx.wrap_socket(raw):
                return "https"
    except Exception:
        return "http"



_AUTOSTART_TASK_NAME = "VigilServeTray"
_AUTOSTART_LEGACY_KEY = r"Software\Microsoft\Windows\CurrentVersion\Run"
_AUTOSTART_LEGACY_VALUE = "VigilServeTray"


def _get_tray_exe_path() -> str:
    """Return the tray executable path."""
    exe = sys.executable
    if not getattr(sys, "frozen", False) and exe.lower().endswith("python.exe"):
        exe = os.path.abspath(sys.argv[0])
    return exe


def _get_tray_task_trigger() -> str:
    """Return the /TR argument quoted for schtasks."""
    exe = _get_tray_exe_path()
    if " " in exe and not exe.startswith('"'):
        exe = f'"{exe}"'
    return exe


def _remove_legacy_autostart() -> None:
    """Remove the old registry Run key if it exists."""
    try:
        with winreg.OpenKey(winreg.HKEY_CURRENT_USER, _AUTOSTART_LEGACY_KEY, 0, winreg.KEY_SET_VALUE) as key:
            try:
                winreg.DeleteValue(key, _AUTOSTART_LEGACY_VALUE)
                logger.info("Removed legacy registry autostart")
            except FileNotFoundError:
                pass
    except Exception as exc:
        logger.warning(f"Failed to remove legacy autostart registry: {exc}")


def _is_autostart_enabled() -> bool:
    """Check whether the tray app is configured to start on boot."""
    try:
        result = subprocess.run(
            ["schtasks", "/Query", "/TN", _AUTOSTART_TASK_NAME, "/FO", "CSV", "/NH"],
            capture_output=True,
            text=True,
            timeout=10,
        )
        if result.returncode != 0:
            return False
        return _get_tray_exe_path() in result.stdout
    except Exception as exc:
        logger.warning(f"Failed to query autostart task: {exc}")
        return False


def _set_autostart(enabled: bool) -> None:
    """Enable or disable tray app autostart via Task Scheduler (highest privileges)."""
    if not getattr(sys, "frozen", False):
        raise RuntimeError("开机自启仅在打包后的 VigilServeTray.exe 中可用")

    _remove_legacy_autostart()

    try:
        if enabled:
            subprocess.run(
                [
                    "schtasks",
                    "/Create",
                    "/F",
                    "/TN", _AUTOSTART_TASK_NAME,
                    "/TR", _get_tray_task_trigger(),
                    "/SC", "ONLOGON",
                    "/RL", "HIGHEST",
                ],
                capture_output=True,
                text=True,
                check=True,
                timeout=30,
            )
            logger.info("Autostart enabled (Task Scheduler, highest privileges)")
        else:
            subprocess.run(
                ["schtasks", "/Delete", "/F", "/TN", _AUTOSTART_TASK_NAME],
                capture_output=True,
                text=True,
                timeout=30,
            )
            logger.info("Autostart disabled")
    except subprocess.CalledProcessError as exc:
        logger.error(f"Failed to update autostart task: {exc.stderr or exc.stdout}")
        raise RuntimeError(f"更新开机自启失败: {exc.stderr or exc.stdout}") from exc
    except Exception as exc:
        logger.error(f"Failed to update autostart task: {exc}")
        raise



def _read_install_options() -> dict:
    """Read install_options.ini written by the installer. Returns {} on missing file.

    Search order (priority high → low):
      1. Next to VigilServeTray.exe
      2. The install root (parent of tray\\), where Inno Setup actually writes
         the file via `ExpandConstant('{app}\\install_options.ini')`
      3. %APPDATA%\\VigilServe\\install_options.ini (per-user override)
    """
    if not getattr(sys, "frozen", False):
        return {}
    exe_dir = os.path.dirname(os.path.abspath(sys.executable))
    appdata = os.environ.get("APPDATA") or os.path.expanduser("~")
    candidates = [
        os.path.join(exe_dir, "install_options.ini"),
        os.path.join(os.path.dirname(exe_dir), "install_options.ini"),
        os.path.join(appdata, "VigilServe", "install_options.ini"),
    ]
    for opts_path in candidates:
        if not os.path.exists(opts_path):
            continue
        out: dict = {}
        try:
            with open(opts_path, "r", encoding="utf-8") as f:
                for raw in f:
                    line = raw.strip()
                    if not line or line.startswith("#") or "=" not in line:
                        continue
                    k, _, v = line.partition("=")
                    out[k.strip()] = v.strip().lower() == "true"
            logger.info(f"Loaded install_options.ini from {opts_path}")
            return out
        except Exception as exc:
            logger.warning(f"Failed to read install_options.ini at {opts_path}: {exc}")
            continue
    logger.info("install_options.ini not found in any candidate location")
    return {}



class ServiceController:
    """Manages the lifecycle of one service."""

    def __init__(self, service: dict):
        self.service = service
        self.log = get_logger(f"service.{service['id']}")
        self.process: subprocess.Popen | None = None
        self._log_file: TextIO | None = None
        self.lock = threading.Lock()
        # must NOT use `self.lock`: the status poller calls `is_running()` from
        # the tkinter thread, so holding `self.lock` across a multi-minute pip
        # install would freeze the whole control panel.
        self._deps_lock = threading.Lock()
        self._deps_ready = False

    def _bootstrap_deps(self) -> None:
        """Run the one-time dependency install, at most once per process."""
        if self._deps_ready:
            return
        with self._deps_lock:
            if self._deps_ready:
                return
            _ensure_service_deps(self.service, self.log)
            self._deps_ready = True

    def is_running(self) -> bool:
        """Return True if the tracked process is alive or the port is occupied."""
        with self.lock:
            if self.process is not None and self.process.poll() is None:
                return True
        port = self.service.get("port")
        if port:
            return _port_in_use(port)
        return False

    def start(self) -> None:
        """Start the service process."""
        if not PYTHON_EXE:
            self.log.error("No usable Python interpreter found; aborting start")
            raise RuntimeError(PYTHON_MISSING_HINT)

        self._bootstrap_deps()
        with self.lock:
            if self.process is not None and self.process.poll() is None:
                return
            cwd = self.service.get("cwd") or os.getcwd()
            if not os.path.isdir(cwd):
                self.log.error(f"Service directory does not exist: {cwd}")
                raise RuntimeError(
                    f"无法启动 {self.service['name']}\n"
                    f"工作目录不存在：{cwd}\n\n"
                    "请确保 VigilServeTray.exe 位于安装目录内运行。"
                )
            env = os.environ.copy()
            env.update(self.service.get("env", {}))
            env["VIGILSERVE_BACKEND_PORT"] = str(self.service.get("port", ""))
            creationflags = subprocess.CREATE_NO_WINDOW
            log_path = _service_log_file(self.service["id"])
            try:
                self._log_file = open(log_path, "a", encoding="utf-8", errors="replace")
                self._log_file.write(f"\n[{datetime.now().isoformat()}] Starting {self.service['name']}\n")
                self._log_file.flush()
                self.process = subprocess.Popen(
                    self.service["cmd"],
                    cwd=cwd,
                    env=env,
                    creationflags=creationflags,
                    stdout=self._log_file,
                    stderr=subprocess.STDOUT,
                )
                self.log.info(f"Started {self.service['name']} pid={self.process.pid}, log={log_path}")
            except Exception as exc:
                self.log.error(f"Failed to start {self.service['name']}: {exc}")
                if self._log_file:
                    self._log_file.close()
                    self._log_file = None
                raise RuntimeError(f"无法启动 {self.service['name']}: {exc}") from exc

            time.sleep(1.0)
            if self.process is not None and self.process.poll() is not None:
                code = self.process.returncode
                self.process = None
                if self._log_file:
                    try:
                        self._log_file.close()
                    except Exception:
                        pass
                    self._log_file = None
                tail = _read_log_tail(log_path)
                self.log.error(
                    f"{self.service['name']} exited immediately (code={code}); log={log_path}"
                )
                hint = f"\n\n{PYTHON_MISSING_HINT}" if not PYTHON_EXE else ""
                raise RuntimeError(
                    f"无法启动 {self.service['name']}：进程启动后立即退出（代码 {code}）。\n\n"
                    f"日志文件：{log_path}\n{tail}{hint}"
                )

    def _kill_by_port(self, port: int) -> None:
        """Find and terminate the process listening on the given port."""
        try:
            for conn in psutil.net_connections(kind="inet"):
                if getattr(conn.laddr, "port", None) == port and conn.pid:
                    try:
                        p = psutil.Process(conn.pid)
                        p.terminate()
                        p.wait(timeout=5)
                        self.log.info(f"Terminated {self.service['name']} on port {port} pid={conn.pid}")
                        return
                    except psutil.NoSuchProcess:
                        pass
                    except subprocess.TimeoutExpired:
                        p.kill()
                        p.wait(timeout=3)
                        self.log.info(f"Killed {self.service['name']} on port {port} pid={conn.pid}")
                        return
        except Exception as exc:
            self.log.error(f"Failed to kill {self.service['name']} by port {port}: {exc}")

    def stop(self) -> None:
        """Stop the tracked process, or any process occupying the service port."""
        with self.lock:
            proc = self.process
            self.process = None
        stopped = False
        if proc is not None and proc.poll() is None:
            try:
                proc.terminate()
                proc.wait(timeout=5)
                self.log.info(f"Stopped {self.service['name']} pid={proc.pid}")
                stopped = True
            except subprocess.TimeoutExpired:
                proc.kill()
                proc.wait(timeout=3)
                self.log.info(f"Killed {self.service['name']} pid={proc.pid}")
                stopped = True
            except Exception as exc:
                self.log.error(f"Error stopping {self.service['name']}: {exc}")
            finally:
                if self._log_file:
                    try:
                        self._log_file.close()
                    except Exception:
                        pass
                    self._log_file = None
        port = self.service.get("port")
        if not stopped and port and _port_in_use(port):
            self._kill_by_port(port)

    def restart(self) -> None:
        """Restart the service."""
        self.stop()
        time.sleep(0.5)
        self.start()



class ControlPanel:
    """A simple control panel window for service status and port config."""

    def __init__(self, app: "TrayApp"):
        self.app = app
        self.root: tk.Tk | None = None
        self._widgets: dict[str, dict] = {}
        self._port_vars: dict[str, tk.StringVar] = {}
        self._old_ports: dict[str, int] = {}
        self._ip_var: tk.StringVar | None = None
        self._logo_img: ImageTk.PhotoImage | None = None
        self._support_images: list[ImageTk.PhotoImage] = []
        self._running = False

    def run(self, visible: bool = True) -> None:
        """Create and run the tkinter main loop in the current thread.

        ``visible=False`` 时窗口建好立刻隐藏。托盘启动时会预建一个隐藏的控制面板
        待命：这样用户之后再双击桌面快捷方式时，第二个进程能用 FindWindow 找到
        这个窗口句柄并激活（见 ``_activate_existing_panel``），不必"先手动打开过
        一次面板"才有窗口可找。
        """
        self.root = tk.Tk()
        self.root.title(f"VigilServeTray - {TRAY_VERSION}")
        self.root.geometry("900x720")
        self.root.resizable(False, False)
        self.root.configure(bg="#ffffff")
        self._set_window_icon()
        self._build_ui()
        self._running = True
        self._refresh_loop()
        self.root.protocol("WM_DELETE_WINDOW", self._on_close)
        if not visible:
            self.root.withdraw()
        self.root.mainloop()
        self._running = False
        self.app._panel = None

    def show(self) -> None:
        """把控制面板窗口显示到前台（已存在时复用，不重建）。"""
        try:
            if not self.root:
                return
            self.root.deiconify()
            self.root.lift()
            self.root.focus_force()
        except Exception as exc:
            logger.warning(f"Show control panel failed: {exc}")

    def _set_window_icon(self) -> None:
        try:
            if os.path.exists(APP_ICON):
                self.root.iconbitmap(APP_ICON)
        except Exception as exc:
            logger.warning(f"Failed to set window icon: {exc}")

    def _build_ui(self) -> None:
        style = ttk.Style(self.root)
        style.configure(
            "TNotebook.Tab",
            font=("Microsoft YaHei", 11),
            padding=(18, 10),
        )
        style.configure(
            "TNotebook",
            tabmargins=(8, 4, 8, 0),
        )
        self._notebook = ttk.Notebook(self.root, style="TNotebook")
        notebook = self._notebook
        notebook.pack(fill=tk.BOTH, expand=True, padx=14, pady=(14, 0))

        service_tab = tk.Frame(notebook, bg="#ffffff")
        notebook.add(service_tab, text="服务配置")
        self._build_service_config_tab(service_tab)

        doc_titles = ["用户协议", "免责声明", "联系作者", "支持作者"]
        for title in doc_titles:
            tab = tk.Frame(notebook, bg="#ffffff")
            notebook.add(tab, text=title)
            if title == "支持作者":
                self._build_support_tab(tab, title)
            else:
                self._build_docx_tab(tab, title)

    def _build_service_config_tab(self, parent: tk.Frame) -> None:
        services_frame = tk.LabelFrame(
            parent,
            text=" 服务状态 ",
            font=("Microsoft YaHei", 11),
            bg="#ffffff",
            fg="#333333",
            bd=1,
            relief=tk.SOLID,
        )
        services_frame.pack(fill=tk.X, padx=20, pady=18)
        for sid, ctrl in self.app.controllers.items():
            row = tk.Frame(services_frame, bg="#ffffff")
            row.pack(fill=tk.X, padx=16, pady=9)
            tk.Label(
                row,
                text=ctrl.service["name"],
                width=16,
                anchor="w",
                bg="#ffffff",
                fg="#333333",
                font=("Microsoft YaHei", 11),
            ).pack(side=tk.LEFT)
            status_label = tk.Label(
                row,
                text="检测中...",
                width=12,
                anchor="center",
                bg="#f5f5f5",
                fg="#666666",
                font=("Microsoft YaHei", 10),
            )
            status_label.pack(side=tk.LEFT, padx=8)
            btn_frame = tk.Frame(row, bg="#ffffff")
            btn_frame.pack(side=tk.RIGHT)
            start_btn = tk.Button(
                btn_frame,
                text="启动",
                width=8,
                font=("Microsoft YaHei", 9),
                command=lambda c=ctrl: self._start(c),
            )
            stop_btn = tk.Button(
                btn_frame,
                text="停止",
                width=8,
                font=("Microsoft YaHei", 9),
                command=lambda c=ctrl: self._stop(c),
            )
            restart_btn = tk.Button(
                btn_frame,
                text="重启",
                width=8,
                font=("Microsoft YaHei", 9),
                command=lambda c=ctrl: self._restart(c),
            )
            for btn in (start_btn, stop_btn, restart_btn):
                btn.pack(side=tk.LEFT, padx=4)
            self._widgets[sid] = {
                "status": status_label,
                "start": start_btn,
                "stop": stop_btn,
                "restart": restart_btn,
            }

        # 服务端 IP：本机网卡地址，供 Agent 端填「服务端地址」（多网卡可下拉选择）
        ip_frame = tk.LabelFrame(
            parent,
            text=" 服务端 IP ",
            font=("Microsoft YaHei", 11),
            bg="#ffffff",
            fg="#333333",
            bd=1,
            relief=tk.SOLID,
        )
        ip_frame.pack(fill=tk.X, padx=20, pady=12)
        ip_row = tk.Frame(ip_frame, bg="#ffffff")
        ip_row.pack(fill=tk.X, padx=16, pady=8)

        _ips = list_local_ips()
        _cur_ip = get_server_ip()
        # 选过但这次没枚举到（比如网线拔了）也要留在下拉里，否则保存的值会被悄悄改掉
        if _cur_ip and _cur_ip not in _ips:
            _ips.insert(0, _cur_ip)
        if not _ips:
            _ips = ["127.0.0.1"]
        self._ip_var = tk.StringVar(value=_cur_ip or _ips[0])

        tk.Label(
            ip_row,
            text="本机地址",
            width=14,
            anchor="w",
            bg="#ffffff",
            fg="#333333",
            font=("Microsoft YaHei", 11),
        ).pack(side=tk.LEFT)
        ip_combo = ttk.Combobox(
            ip_row,
            textvariable=self._ip_var,
            values=_ips,
            state="readonly",
            width=22,
            font=("Microsoft YaHei", 11),
        )
        ip_combo.pack(side=tk.LEFT)
        # 选完立刻落盘：这个选项只影响"显示与复制地址"，没有需要重启的服务，
        # 不需要再多一个"保存"按钮打断用户。
        ip_combo.bind("<<ComboboxSelected>>", lambda _e: set_server_ip(self._ip_var.get()))
        tk.Label(
            ip_frame,
            text="Agent 端「服务端地址」请填这个 IP（多网卡时选 Agent 能访问到的那张网卡）",
            anchor="w",
            bg="#ffffff",
            fg="#888888",
            font=("Microsoft YaHei", 9),
        ).pack(fill=tk.X, padx=16, pady=(0, 8))

        ports_frame = tk.LabelFrame(
            parent,
            text=" 端口配置 ",
            font=("Microsoft YaHei", 11),
            bg="#ffffff",
            fg="#333333",
            bd=1,
            relief=tk.SOLID,
        )
        ports_frame.pack(fill=tk.X, padx=20, pady=12)
        ports = get_ports()
        self._old_ports = dict(ports)
        labels = [("后端端口", "backend"), ("Agent 端口", "agent"), ("WEB代理端口", "webproxy")]
        for label_text, key in labels:
            row = tk.Frame(ports_frame, bg="#ffffff")
            row.pack(fill=tk.X, padx=16, pady=8)
            tk.Label(
                row,
                text=label_text,
                width=14,
                anchor="w",
                bg="#ffffff",
                fg="#333333",
                font=("Microsoft YaHei", 11),
            ).pack(side=tk.LEFT)
            var = tk.StringVar(value=str(ports[key]))
            entry = tk.Entry(row, textvariable=var, width=14, font=("Microsoft YaHei", 11))
            entry.pack(side=tk.LEFT)
            self._port_vars[key] = var

        tk.Button(
            ports_frame,
            text="保存端口",
            command=self._save_ports,
            font=("Microsoft YaHei", 10),
            width=12,
        ).pack(anchor="e", padx=16, pady=8)

        footer = tk.Frame(parent, bg="#ffffff")
        footer.pack(fill=tk.X, padx=20, pady=18)
        tk.Button(
            footer,
            text="打开 Web 控制台",
            command=self._open_web,
            font=("Microsoft YaHei", 10),
            width=16,
        ).pack(side=tk.LEFT, padx=4)
        tk.Button(
            footer,
            text="刷新状态",
            command=self._refresh,
            font=("Microsoft YaHei", 10),
            width=12,
        ).pack(side=tk.RIGHT, padx=4)

    def _build_docx_tab(self, parent: tk.Frame, title: str) -> None:
        """Build a read-only text tab from the embedded docs bundle."""
        text_widget = tk.Text(
            parent,
            wrap=tk.WORD,
            font=("Microsoft YaHei", 11),
            padx=16,
            pady=16,
            bg="#ffffff",
            fg="#333333",
            relief=tk.FLAT,
            highlightthickness=0,
        )
        scrollbar = tk.Scrollbar(parent, command=text_widget.yview)
        text_widget.configure(yscrollcommand=scrollbar.set)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        text_widget.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=10, pady=10)

        blocks = self._load_doc_content(title)
        self._render_blocks(text_widget, blocks)
        text_widget.config(state=tk.DISABLED)

    def _load_doc_content(self, title: str) -> list[dict]:
        """Load document blocks from the embedded bundle, falling back to .docx."""
        if title in DOCS_BUNDLE:
            return DOCS_BUNDLE[title].get("blocks", [])

        docx_path = _docx_path(_docx_filename(title))
        if Document is None:
            return [{"type": "paragraph", "text": f"无法读取文档：缺少 python-docx 依赖\n{docx_path}"}]
        try:
            doc = Document(docx_path)
            blocks = []
            for para in doc.paragraphs:
                if not para.text.strip() and not para.runs:
                    continue
                blocks.append(_extract_paragraph_block(para))
            for table in doc.tables:
                blocks.append(_extract_table_block(table))
            return blocks
        except Exception as exc:
            return [{"type": "paragraph", "text": f"无法读取文档：{exc}\n{docx_path}"}]

    def _render_blocks(self, text_widget: tk.Text, blocks: list[dict]) -> None:
        """Render structured docx blocks into a tkinter Text widget."""
        text_widget.tag_configure("bold", font=("Microsoft YaHei", 11, "bold"))
        text_widget.tag_configure("italic", font=("Microsoft YaHei", 11, "italic"))
        text_widget.tag_configure("center", justify="center")
        text_widget.tag_configure("right", justify="right")
        text_widget.tag_configure("table", font=("Consolas", 10))

        for idx, block in enumerate(blocks):
            btype = block.get("type", "paragraph")
            if btype == "paragraph":
                self._render_paragraph(text_widget, block)
            elif btype == "table":
                self._render_table(text_widget, block)
            if btype != "paragraph" or block.get("text", "").strip():
                text_widget.insert(tk.END, "\n")

    def _render_paragraph(self, text_widget: tk.Text, block: dict) -> None:
        """Insert a paragraph with formatting tags."""
        text = block.get("text", "")
        align = block.get("align")
        left_indent = block.get("left_indent") or 0
        first_indent = block.get("first_line_indent") or 0

        tags: list[str] = []
        if align == "center":
            tags.append("center")
        elif align == "right":
            tags.append("right")

        lmargin1 = int((left_indent + first_indent) * 1.333) if (left_indent or first_indent) else 0
        lmargin2 = int(left_indent * 1.333) if left_indent else 0
        if lmargin1 or lmargin2:
            tag_name = f"indent_{int(left_indent)}_{int(first_indent)}"
            text_widget.tag_configure(tag_name, lmargin1=lmargin1, lmargin2=lmargin2)
            tags.append(tag_name)

        runs = block.get("runs") or [{"text": text, "bold": block.get("bold", False), "italic": False, "underline": False}]
        for run in runs:
            run_tags = list(tags)
            if run.get("bold"):
                run_tags.append("bold")
            if run.get("italic"):
                run_tags.append("italic")
            if run_tags:
                text_widget.insert(tk.END, run["text"], tuple(run_tags))
            else:
                text_widget.insert(tk.END, run["text"])
        text_widget.insert(tk.END, "\n", tuple(tags) if tags else ())

    def _render_table(self, text_widget: tk.Text, block: dict) -> None:
        """Render a table as a ttk.Treeview embedded in the Text widget."""
        rows = block.get("rows", [])
        if not rows:
            return

        container = tk.Frame(text_widget, bg="#ffffff")
        tree = ttk.Treeview(
            container,
            columns=list(range(len(rows[0]))),
            show="headings",
            height=min(len(rows), 15),
        )
        tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

        for ci, header in enumerate(rows[0]):
            tree.heading(ci, text=header)
            max_len = max(len(str(row[ci])) if ci < len(row) else 0 for row in rows)
            display_len = max(len(str(header)), max_len)
            tree.column(ci, width=min(display_len * 14 + 20, 300), anchor="w")

        for row in rows[1:]:
            tree.insert("", tk.END, values=row)

        text_widget.insert(tk.END, "\n")
        text_widget.window_create(tk.END, window=container)
        text_widget.insert(tk.END, "\n")


    def _build_support_tab(self, parent: tk.Frame, title: str) -> None:
        """Build the '支持作者' tab: intro text on top, donation cards below."""
        blocks, images = _load_support_content(title)
        _render_support_content(parent, blocks, images, self._support_images)

    def _refresh(self) -> None:
        try:
            for sid, ctrl in self.app.controllers.items():
                running = ctrl.is_running()
                self._set_status(sid, running)
        except Exception as exc:
            logger.error(f"Refresh status failed: {exc}")

    def _set_status(self, sid: str, running: bool) -> None:
        widgets = self._widgets.get(sid)
        if not widgets:
            return
        if running:
            widgets["status"].config(text="运行中", bg="#e6f7e6", fg="#237804")
        else:
            widgets["status"].config(text="已停止", bg="#fff2f0", fg="#cf1322")
        widgets["start"].config(state=tk.NORMAL if not running else tk.DISABLED)
        widgets["stop"].config(state=tk.NORMAL if running else tk.DISABLED)
        widgets["restart"].config(state=tk.NORMAL if running else tk.DISABLED)

    def _refresh_loop(self) -> None:
        if not self._running or not self.root:
            return
        # 隐藏时跳过刷新：控制面板窗口在托盘启动时就预建好了（隐藏待命），
        # 不判断的话每秒一次的状态轮询会在后台白跑一整天。
        try:
            if self.root.state() != "withdrawn":
                self._refresh()
        except Exception:
            self._refresh()
        self.root.after(1000, self._refresh_loop)

    def _run_safe(self, fn) -> None:
        try:
            fn()
        except Exception as exc:
            # 🚨 2026-09-23 修复：下面弹窗用的是 `root.after(0, ...)` —— 它是**延迟到
            # 主线程下一轮**才执行的，而 Python 3 在 `except` 块结束时会自动 `del exc`。
            # 原来写 `lambda: messagebox.showerror(..., str(exc))`，闭包捕获的是这个名字
            # 本身，等到回调真正跑起来时 `exc` 已经被删掉了 → 弹窗的瞬间再抛
            # `NameError: name 'exc' is not defined`，又被最里面那句裸 `except` 吞掉。
            # 用户看到的现象：点了「启动/停止」没反应，**连一个错误提示都没有**。
            #
            # 所以先把文本取出来，再用默认参数（`lambda m=msg:`）在**定义时**绑定，
            # 与这个文件里其它几处 `lambda c=ctrl:` 的写法保持一致。
            msg = str(exc) or exc.__class__.__name__
            logger.error(f"Control panel action failed: {msg}")
            if self.root:
                try:
                    self.root.after(
                        0, lambda m=msg: messagebox.showerror("操作失败", m, parent=self.root))
                except Exception as _box_err:  # noqa: BLE001
                    # 弹窗本身也可能失败（窗口已销毁等）。以前这里静默吞掉，正是它把
                    # 上面那个 NameError 藏了一年多——失败至少要在日志里留痕。
                    logger.error(f"显示错误弹窗失败：{_box_err!r}")

    def _start(self, ctrl: ServiceController) -> None:
        threading.Thread(target=self._run_safe, args=(ctrl.start,), daemon=True).start()

    def _stop(self, ctrl: ServiceController) -> None:
        threading.Thread(target=self._run_safe, args=(ctrl.stop,), daemon=True).start()

    def _restart(self, ctrl: ServiceController) -> None:
        threading.Thread(target=self._run_safe, args=(ctrl.restart,), daemon=True).start()

    def _save_ports(self) -> None:
        try:
            new_ports = {k: int(v.get()) for k, v in self._port_vars.items()}
            for p in new_ports.values():
                if not (1 <= p <= 65535):
                    raise ValueError("端口必须在 1-65535 之间")
        except ValueError as exc:
            messagebox.showerror("输入错误", f"端口格式无效：{exc}", parent=self.root)
            return

        changed = [k for k in new_ports if new_ports[k] != self._old_ports.get(k)]
        set_ports(new_ports)
        self.app.reload_controllers()
        self._old_ports = dict(new_ports)

        if changed and messagebox.askyesno(
            "重启服务",
            f"以下服务端口已更改：{', '.join(changed)}\n是否立即重启这些服务？",
            parent=self.root,
        ):
            for k in changed:
                ctrl = self.app.controllers.get(k)
                if ctrl:
                    threading.Thread(target=ctrl.restart, daemon=True).start()

        messagebox.showinfo("保存成功", "端口配置已保存", parent=self.root)

    def _open_web(self) -> None:
        self.app._open_console(None, None)

    def _on_close(self) -> None:
        # 关闭 = 隐藏，不销毁：窗口句柄要一直留着，第二个进程（用户双击桌面
        # 快捷方式）要靠 FindWindow 找到它再激活。销毁掉的话下次双击就只能弹
        # "已在运行"，用户还是看不到面板。托盘图标本身不受影响，照旧驻留。
        try:
            self.root.withdraw()
        except Exception as exc:
            logger.warning(f"Hide control panel failed: {exc}")



class TrayApp:
    """System tray application."""

    def __init__(self):
        logger.info("TrayApp initializing")
        self.controllers: dict[str, ServiceController] = {}
        self.icon: pystray.Icon | None = None
        self._icon_image = self._load_icon()
        self._autostart_enabled = _is_autostart_enabled()
        self._install_options = _read_install_options()
        self._controllers_lock = threading.Lock()
        self._panel: ControlPanel | None = None
        self._support_window: tk.Toplevel | None = None
        self._support_image_refs: list[ImageTk.PhotoImage] = []
        self._rebuild_controllers()
        logger.info("TrayApp initialized")
        if self._install_options.get("launch_service", False):
            threading.Timer(0.8, self._start_all_services_async).start()

    def _rebuild_controllers(self) -> None:
        """Rebuild service controllers from current config."""
        with self._controllers_lock:
            old = dict(self.controllers)
            self.controllers = {}
            for svc in SERVICES:
                sid = svc["id"]
                if sid in old and old[sid].service["port"] == svc["port"]:
                    self.controllers[sid] = old[sid]
                else:
                    self.controllers[sid] = ServiceController(svc)

    def reload_controllers(self) -> None:
        """Reload config and rebuild controllers (called after port change)."""
        global SERVICES
        SERVICES[:] = build_services()
        self._rebuild_controllers()
        self._update_tooltip()

    def _load_icon(self):
        icon_path = APP_ICON
        if not os.path.exists(icon_path):
            logger.warning(f"Application icon not found at {icon_path}, using fallback")
            return Image.new("RGBA", (64, 64), (0, 120, 215, 255))
        try:
            img = Image.open(icon_path)
            img = img.convert("RGBA")
            img.thumbnail((64, 64))
            logger.info(f"Loaded icon from {icon_path}")
            return img
        except Exception as exc:
            logger.error(f"Failed to load icon: {exc}")
            return Image.new("RGBA", (64, 64), (0, 120, 215, 255))

    def _status_text(self) -> str:
        lines = []
        for sid, ctrl in self.controllers.items():
            state = "运行中" if ctrl.is_running() else "已停止"
            lines.append(f"{ctrl.service['name']}: {state}")
        return "\n".join(lines) or "VigilServe 服务控制台"

    def _ensure_panel(self, visible: bool = True) -> None:
        """保证控制面板窗口存在；``visible=True`` 时把它显示到前台。

        窗口是在单独的线程里跑 tkinter mainloop 的，所以"确保存在"是异步的：
        这里只是把建窗线程拉起来，真正可交互要等那一小会儿。
        """
        if self._panel is not None:
            if visible:
                self._panel.show()
            return

        def panel_thread():
            try:
                panel = ControlPanel(self)
                self._panel = panel
                panel.run(visible=visible)
            except Exception as exc:
                logger.exception(f"Control panel failed: {exc}")
                self._panel = None

        threading.Thread(target=panel_thread, daemon=True).start()

    def _open_panel(self, icon, item):
        """Open the tkinter control panel."""
        logger.info("Open control panel selected")
        self._ensure_panel(visible=True)

    def _open_support_author(self, icon, item):
        """Open a standalone support-author window from the tray menu."""
        logger.info("Open support author selected")
        if self._support_window is not None:
            try:
                self._support_window.lift()
                self._support_window.focus_force()
            except Exception:
                pass
            return

        def support_thread():
            try:
                root = tk.Tk()
                root.withdraw()
                window = tk.Toplevel(root)
                window.title("支持作者")
                window.geometry("880x620")
                window.resizable(False, False)
                window.configure(bg="#ffffff")
                try:
                    if os.path.exists(APP_ICON):
                        window.iconbitmap(APP_ICON)
                except Exception as exc:
                    logger.warning(f"Failed to set support window icon: {exc}")

                blocks, images = _load_support_content("支持作者")
                _render_support_content(window, blocks, images, self._support_image_refs, wraplength=740)

                self._support_window = window

                def _on_close():
                    self._support_window = None
                    self._support_image_refs.clear()
                    try:
                        window.destroy()
                    except Exception:
                        pass
                    try:
                        root.destroy()
                    except Exception:
                        pass

                window.protocol("WM_DELETE_WINDOW", _on_close)
                window.mainloop()
            except Exception as exc:
                logger.exception(f"Support author window failed: {exc}")
                self._support_window = None

        threading.Thread(target=support_thread, daemon=True).start()

    def _open_console(self, icon, item):
        logger.info("Open web console selected")
        try:
            ports = get_ports()
            port = ports.get("backend", 8001)
            scheme = _backend_scheme(port)
            webbrowser.open(f"{scheme}://127.0.0.1:{port}")
        except Exception as exc:
            logger.error(f"Open console failed: {exc}")

    def _restart_backend(self, icon, item):
        logger.info("Restart backend selected")
        ctrl = self.controllers.get("backend")
        if ctrl:
            threading.Thread(target=ctrl.restart, daemon=True).start()

    def _start_all_services_async(self) -> None:
        """Auto-start all services (used after install when launch_service=true)."""
        logger.info("Auto-starting services (install_options.ini: launch_service=true)")
        for ctrl in self.controllers.values():
            try:
                if not ctrl.is_running():
                    threading.Thread(target=ctrl.start, daemon=True).start()
            except Exception as exc:
                logger.error(f"Auto-start failed for {ctrl.service['name']}: {exc}")
        self._update_tooltip()

    def _toggle_autostart(self, icon, item):
        try:
            self._autostart_enabled = not self._autostart_enabled
            _set_autostart(self._autostart_enabled)
            if self.icon:
                self.icon.update_menu()
        except Exception as exc:
            logger.error(f"Toggle autostart failed: {exc}")
            self._autostart_enabled = not self._autostart_enabled

    def _autostart_checked(self, item):
        return self._autostart_enabled

    def _exit(self, icon, item):
        logger.info("Exit selected from tray menu")
        try:
            if self._panel and self._panel.root:
                self._panel._on_close()
        except Exception:
            pass
        try:
            icon.stop()
        except Exception as exc:
            logger.error(f"icon.stop failed: {exc}")
        for ctrl in self.controllers.values():
            ctrl.stop()
        os._exit(0)

    def _build_menu(self):
        return pystray.Menu(
            pystray.MenuItem("打开控制台", self._open_panel, default=True),
            pystray.Menu.SEPARATOR,
            pystray.MenuItem("打开 Web 控制台", self._open_console),
            pystray.Menu.SEPARATOR,
            pystray.MenuItem("重启后端服务", self._restart_backend),
            pystray.Menu.SEPARATOR,
            pystray.MenuItem("开机自启", self._toggle_autostart, checked=self._autostart_checked),
            pystray.Menu.SEPARATOR,
            pystray.MenuItem("支持作者", self._open_support_author),
            pystray.Menu.SEPARATOR,
            pystray.MenuItem("退出", self._exit),
        )

    def _update_tooltip(self):
        if self.icon:
            try:
                self.icon.title = self._status_text()
            except Exception as exc:
                logger.warning(f"Update tooltip failed: {exc}")

    def _tooltip_worker(self):
        while True:
            try:
                self._update_tooltip()
            except Exception as exc:
                logger.error(f"Tooltip worker failed: {exc}")
            time.sleep(3)

    def run(self) -> None:
        threading.Thread(target=self._tooltip_worker, daemon=True).start()
        menu = self._build_menu()
        self.icon = pystray.Icon(
            "VigilServeTray",
            self._icon_image,
            self._status_text(),
            menu,
        )
        logger.info("Tray icon starting")
        # 预建控制面板窗口：带 --panel 就直接显示（装完 / 双击快捷方式），
        # 否则隐藏待命，等第二个进程来唤醒。建窗失败不影响托盘本身。
        try:
            self._ensure_panel(visible=_wants_panel())
        except Exception as exc:
            logger.warning(f"Pre-create control panel failed: {exc}")
        self.icon.run()
        logger.info("Tray icon stopped")


def _acquire_single_instance() -> bool:
    """Ensure only one tray instance runs (named-mutex guard).

    Double-clicking the desktop shortcut while the tray is already running
    used to spawn a second tray icon / control panel. CreateMutexW with a
    fixed name makes the second process detect the first one and exit.
    Returns True when this process owns the single-instance mutex.
    """
    if sys.platform != "win32":
        return True
    try:
        import ctypes

        ERROR_ALREADY_EXISTS = 183
        handle = ctypes.windll.kernel32.CreateMutexW(
            None, False, "VigilServeTray-SingleInstance"
        )
        if ctypes.windll.kernel32.GetLastError() == ERROR_ALREADY_EXISTS:
            return False
        globals()["_SINGLE_INSTANCE_MUTEX"] = handle
        return True
    except Exception as exc:
        logger.warning(f"Single-instance check failed: {exc}")
        return True


def _wants_panel() -> bool:
    """命令行是否带了 ``--panel``：启动后直接把控制面板窗口弹出来。

    桌面 / 开始菜单快捷方式都带这个参数（见 setup.iss 的 ``Parameters``），
    开机自启项**不带** —— 否则每次登录都弹窗，很烦。
    """
    try:
        return any(str(a).strip().lower() == "--panel" for a in sys.argv[1:])
    except Exception:
        return False


def _activate_existing_panel() -> bool:
    """把已经在运行的那个实例的控制面板窗口唤到前台。

    托盘是单实例的：第二个进程拿不到互斥体，只能退出。以前它只会弹一句
    "已在运行"，用户双击桌面快捷方式什么也看不到。现在按窗口标题找到第一个
    实例预建的面板窗口并激活 —— 窗口由 ``TrayApp.run`` 的预建逻辑保证一定存在
    （隐藏的也算），确实找不到就返回 False，由调用方退回提示框。
    """
    if sys.platform != "win32":
        return False
    try:
        import ctypes

        user32 = ctypes.windll.user32
        user32.FindWindowW.restype = ctypes.c_void_p
        user32.FindWindowW.argtypes = [ctypes.c_wchar_p, ctypes.c_wchar_p]
        user32.ShowWindow.argtypes = [ctypes.c_void_p, ctypes.c_int]
        user32.SetForegroundWindow.argtypes = [ctypes.c_void_p]
        hwnd = user32.FindWindowW(None, f"VigilServeTray - {TRAY_VERSION}")
        if not hwnd:
            return False
        SW_RESTORE = 9
        user32.ShowWindow(hwnd, SW_RESTORE)
        user32.SetForegroundWindow(hwnd)
        return True
    except Exception as exc:
        logger.warning(f"Activate existing control panel failed: {exc}")
        return False


def main():
    if not _acquire_single_instance():
        logger.info("Another VigilServeTray instance is already running, exiting")
        # 带 --panel（双击桌面快捷方式）时直接把已有实例的面板唤出来，
        # 别再弹个"已在运行"就把用户打发了。
        if _wants_panel() and _activate_existing_panel():
            return
        try:
            import ctypes

            ctypes.windll.user32.MessageBoxW(
                0, "VigilServe 托盘程序已在运行。\n\n"
                   "可双击系统托盘图标、或右键选择「打开控制台」。",
                "VigilServe", 0x40
            )
        except Exception:
            pass
        return
    if sys.platform == "win32" and hasattr(sys, "frozen"):
        try:
            import ctypes
            ctypes.windll.user32.ShowWindow(ctypes.windll.kernel32.GetConsoleWindow(), 0)
        except Exception as exc:
            logger.warning(f"Hide console window failed: {exc}")
    try:
        app = TrayApp()
        app.run()
    except Exception as exc:
        logger.exception(f"Fatal error: {exc}")
        raise


if __name__ == "__main__":
    main()
