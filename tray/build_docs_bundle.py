"""Build an embedded docs bundle from the four .docx files in tray/docs/.

Run this script after editing the .docx files to regenerate tray/docs_bundle.py.
The generated module is imported by tray_app.py at runtime, so the final
VigilServeTray.exe contains the document text, tables and images as Python
bytecode instead of shipping loose .docx files that end users can edit.
"""

from __future__ import annotations

import base64
import json
import re
import sys
import textwrap
from pathlib import Path
from typing import Any

TRAY_DIR = Path(__file__).parent.resolve()
sys.path.insert(0, str(TRAY_DIR))

from docx import Document
from docx.oxml.ns import qn


def _extract_inline_images(doc: Document) -> list[dict[str, Any]]:
    """Return inline images in document order as {ext, data_base64} dicts."""
    images: list[dict[str, Any]] = []
    seen: set[str] = set()
    for para in doc.paragraphs:
        for run in para._p.iter():
            if run.tag.endswith("}drawing") or run.tag.endswith("drawing"):
                for blip in run.iter():
                    if blip.tag.endswith("}blip") or blip.tag.endswith("blip"):
                        embed = blip.get(qn("r:embed"))
                        if not embed or embed in seen:
                            continue
                        seen.add(embed)
                        try:
                            image_part = doc.part.related_parts[embed]
                            blob = image_part.blob
                            ext = _guess_ext(image_part.content_type)
                            images.append(
                                {
                                    "ext": ext,
                                    "data_base64": base64.b64encode(blob).decode("ascii"),
                                }
                            )
                        except Exception as exc:
                            print(f"Warning: failed to extract image {embed}: {exc}")
    return images


def _guess_ext(content_type: str) -> str:
    mapping = {
        "image/png": ".png",
        "image/jpeg": ".jpg",
        "image/gif": ".gif",
        "image/bmp": ".bmp",
    }
    return mapping.get(content_type.lower(), ".bin")


def _pt(value: Any) -> float | None:
    """Convert a length object to points, or return None."""
    if value is None:
        return None
    try:
        return value.pt
    except Exception:
        return None


def _align_name(value: Any) -> str | None:
    """Convert docx paragraph alignment to a simple string."""
    if value is None:
        return None
    mapping = {0: "left", 1: "center", 2: "right", 3: "justify"}
    return mapping.get(value, "left")


def _extract_paragraph(para) -> dict[str, Any]:
    """Extract paragraph text with inline formatting preserved per run."""
    runs = []
    for run in para.runs:
        text = run.text
        if not text:
            continue
        runs.append(
            {
                "text": text,
                "bold": bool(run.bold),
                "italic": bool(run.italic),
                "underline": bool(run.underline),
            }
        )
    return {
        "type": "paragraph",
        "text": para.text,
        "runs": runs,
        "bold": bool(para.runs and all(r.bold for r in para.runs if r.text)),
        "align": _align_name(para.alignment),
        "first_line_indent": _pt(para.paragraph_format.first_line_indent),
        "left_indent": _pt(para.paragraph_format.left_indent),
        "style": para.style.name if para.style else None,
    }


def _extract_table(table) -> dict[str, Any]:
    """Extract a docx table as a list of rows."""
    rows = []
    for row in table.rows:
        rows.append([cell.text.strip() for cell in row.cells])
    return {"type": "table", "rows": rows}


def build_bundle() -> dict[str, Any]:
    docs_dir = TRAY_DIR / "docs"
    files = {
        "用户协议": "VigilServe用户协议.docx",
        "免责声明": "VigilServe免责声明.docx",
        "联系作者": "联系作者.docx",
        "支持作者": "支持作者.docx",
    }
    bundle: dict[str, Any] = {}
    for title, filename in files.items():
        path = docs_dir / filename
        if not path.exists():
            raise FileNotFoundError(f"Required docx not found: {path}")
        doc = Document(path)
        blocks: list[dict[str, Any]] = []
        for child in doc.element.body.iterchildren():
            tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
            if tag == "p":
                para = doc.paragraphs[0].__class__(child, doc.part)
                if not para.text.strip() and not para.runs:
                    continue
                blocks.append(_extract_paragraph(para))
            elif tag == "tbl":
                table = doc.tables[0].__class__(child, doc.part)
                blocks.append(_extract_table(table))
        bundle[title] = {
            "blocks": blocks,
            "images": _extract_inline_images(doc),
        }
    return bundle


def _format_bundle_source(bundle: dict[str, Any]) -> str:
    """Produce a self-contained Python module string."""
    payload = json.dumps(bundle, ensure_ascii=False, indent=2)
    payload = re.sub(r"\btrue\b", "True", payload)
    payload = re.sub(r"\bfalse\b", "False", payload)
    payload = re.sub(r"\bnull\b", "None", payload)
    indented_payload = textwrap.indent(payload, "    ")
    return (
        '"""Embedded document bundle for VigilServe Tray.\n\n'
        "This file is auto-generated by build_docs_bundle.py. Do not edit by hand.\n"
        "Run ``python build_docs_bundle.py`` from the tray directory after\n"
        "modifying the .docx source files.\n"
        '"""\n\n'
        "# ruff: noqa\n"
        "# fmt: off\n"
        f"DOCS_BUNDLE = {indented_payload.lstrip()}\n\n"
        'CAPTIONS = ["支付宝", "微信支付"]\n'
    )


def main() -> None:
    bundle = build_bundle()
    output_path = TRAY_DIR / "docs_bundle.py"
    output_path.write_text(_format_bundle_source(bundle), encoding="utf-8")
    print(f"Generated embedded docs bundle: {output_path}")
    for title, data in bundle.items():
        img_count = len(data["images"])
        block_count = len(data["blocks"])
        print(f"  {title}: {block_count} block(s), {img_count} image(s)")


if __name__ == "__main__":
    main()
