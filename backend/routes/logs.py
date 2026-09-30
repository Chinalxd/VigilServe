"""Log management API: alert logs and operation logs."""
import io
import re
from datetime import datetime, timezone, timedelta
from typing import Optional
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Header, Query, Request
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy import or_
from sqlalchemy.orm import Session

from database import get_db
from models import Alert, OperationLog, Server, User
from routes.auth import (require_admin_full, require_login,
                         require_perm, allowed_server_ids, can_manage_server)
from services.audit_logger import log_operation
from services.client_ip import get_client_ip

router = APIRouter(prefix="/api/logs", tags=["logs"])


def _require_logs(authorization, db: Session, action: str = "view") -> dict:
    """日志接口统一门禁（2026-09-22 安全审计 P1-2）。

    修复前这 6 个端点只有 `require_login`，**不校验权限点**：前端按钮受
    `sys.logs.view` / `sys.logs.export` 控制，后端却来者不拒 ——
    只要绕过前端直接调接口，任何已登录账号都能导出全量审计日志。
    现在按操作类型区分：查看 = view，导出 = export。
    """
    user = require_login(authorization, db)
    require_perm(db, user, "sys", "logs", action,
                 detail="无权限查看日志" if action == "view" else "无权限导出日志")
    return user


def _log_scope(db: Session, user: dict, server_id: Optional[int]) -> list:
    """返回当前用户在日志里可见的主机 ID 列表；指定越权主机时直接 403。

    管理员 = 全部；其余角色 = 角色「管理主机」列表。列表为空 → 看不到任何主机日志。
    （注意：只收敛**主机维度**的日志。操作日志里还有非主机类的记录，
     那部分不在此收敛范围内，避免过度收紧把正常审计视图搞坏。）
    """
    all_ids = [r[0] for r in db.query(Server.id).all()]
    allowed = allowed_server_ids(db, user, all_ids)
    if server_id is not None:
        if not can_manage_server(db, user, server_id):
            raise HTTPException(status_code=403, detail="该主机不在当前角色的管理范围内")
    return allowed

DEFAULT_PAGE_SIZE = 50
CONTEXT_WINDOW = 5
ACTIVE_LOG_PROTECTION_HOURS = 1
CLEAR_MAX_LIMIT = 100000


def _client_ip(request: Request) -> str:
    """⚠ 2026-09-23 动态审计 D-1：取值规则已统一到 `services/client_ip.py`
    （默认不信任 X-Forwarded-For，只认直连 IP；确实走了反向代理时用
    `VIGILSERVE_TRUSTED_PROXIES` 显式声明代理地址）。这里保留同名薄封装，
    是为了不动本文件里散落的 `ip_address=_client_ip(request)`。"""
    return get_client_ip(request)


def _parse_dt(value: Optional[str]) -> Optional[datetime]:
    if not value:
        return None
    try:
        dt = datetime.fromisoformat(value.replace("Z", "+00:00"))
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        return dt
    except Exception:
        return None


def _ts(dt: datetime) -> str:
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.isoformat()


def _filename_dt(value: Optional[str]) -> str:
    """Format an ISO datetime string as 'YYYY-MM-DD HH:MM' for filenames."""
    dt = _parse_dt(value)
    if not dt:
        return ""
    return dt.strftime("%Y-%m-%d %H:%M")


def _apply_keyword_filter(query, model, keyword: Optional[str], regex: bool):
    if not keyword:
        return query, None, None
    fields = []
    if hasattr(model, "message"):
        fields.append(model.message)
    if hasattr(model, "title"):
        fields.append(model.title)
    if hasattr(model, "username"):
        fields.append(model.username)
    if hasattr(model, "action"):
        fields.append(model.action)
    if hasattr(model, "category"):
        fields.append(model.category)
    if hasattr(model, "target_type"):
        fields.append(model.target_type)
    if hasattr(model, "target_id"):
        fields.append(model.target_id)

    if regex:
        return query, keyword, fields
    pattern = f"%{keyword}%"
    clauses = [f.ilike(pattern) for f in fields]
    return query.filter(or_(*clauses)), None, None


def _regex_filter_rows(rows, keyword: str, fields):
    try:
        rx = re.compile(keyword, re.IGNORECASE)
    except re.error as e:
        raise HTTPException(status_code=400, detail=f"正则表达式错误: {e}")

    def match(row):
        for f in fields:
            val = getattr(row, f.key, None)
            if val is not None and rx.search(str(val)):
                return True
        return False

    return [r for r in rows if match(r)]


def _alert_base_query(db: Session, start: Optional[str], end: Optional[str],
                      server_id: Optional[int], level: Optional[str],
                      allowed_ids: Optional[list] = None):
    query = db.query(Alert, Server.name.label("server_name")).join(
        Server, Alert.server_id == Server.id, isouter=True
    )
    # 2026-09-22 P1-2：非管理员只看到自己管理范围内的主机告警。
    # 管理员不传这个参数（可见全部），所以行为完全不变。
    if allowed_ids is not None:
        query = query.filter(Alert.server_id.in_(allowed_ids))
    t0 = _parse_dt(start)
    t1 = _parse_dt(end)
    if t0:
        query = query.filter(Alert.timestamp >= t0)
    if t1:
        query = query.filter(Alert.timestamp <= t1)
    if server_id is not None:
        query = query.filter(Alert.server_id == server_id)
    if level:
        query = query.filter(Alert.level == level)
    return query


def _serialize_alert(row, server_name: str = ""):
    """把一行告警序列化成 dict。

    🚨 别用 ``isinstance(row, tuple)`` 判断"这是不是 (Alert, server_name) 这种
    join 查询行"：SQLAlchemy 2.x 的 ``Row`` **不是** ``tuple`` 的子类，判断会静默
    变成 False，于是整行被当成 Alert 使用 → ``row.id`` 抛
    ``AttributeError: id``（BaseRow._key_not_found）→ /api/logs/alerts 直接 500，
    前端的报错文案就是 "Failed to fetch alert logs"。
    改成：先认 Alert 实例，其余一律按"可按位置取值"的行来处理。
    """
    if isinstance(row, Alert):
        alert, sname = row, server_name
    else:
        vals = tuple(row)
        alert = vals[0]
        sname = server_name or (vals[1] if len(vals) > 1 else "")
    return {
        "id": alert.id,
        "server_id": alert.server_id,
        "server_name": sname or "",
        "timestamp": _ts(alert.timestamp),
        "level": alert.level,
        "title": alert.title,
        "message": alert.message,
        "metric_type": alert.metric_type,
        "acknowledged": alert.acknowledged,
    }


def _op_base_query(db: Session, start: Optional[str], end: Optional[str],
                   category: Optional[str], username: Optional[str],
                   server_id: Optional[int] = None):
    query = db.query(OperationLog)
    # 主机详情页「事件日志 / 操作日志」页签：只保留与这台主机相关的记录
    if server_id is not None:
        query = query.filter(
            OperationLog.target_type == "server",
            OperationLog.target_id == str(server_id),
        )
    t0 = _parse_dt(start)
    t1 = _parse_dt(end)
    if t0:
        query = query.filter(OperationLog.timestamp >= t0)
    if t1:
        query = query.filter(OperationLog.timestamp <= t1)
    if category:
        query = query.filter(OperationLog.category == category)
    if username:
        matched_logins = (
            db.query(User.username)
            .filter(User.full_name.ilike(f"%{username}%"))
            .subquery()
        )
        query = query.filter(
            or_(
                OperationLog.username.ilike(f"%{username}%"),
                OperationLog.username.in_(matched_logins),
            )
        )
    return query


def _build_user_map(db: Session, logs: list[OperationLog]) -> dict[str, str]:
    """Build a username -> full_name map for the given log records."""
    usernames = {log.username for log in logs if log.username}
    if not usernames:
        return {}
    return {
        u.username: (u.full_name or "")
        for u in db.query(User).filter(User.username.in_(usernames)).all()
    }


def _serialize_op(log: OperationLog, user_map: dict[str, str] | None = None):
    full_name = ""
    if user_map is not None:
        full_name = user_map.get(log.username, "")
    return {
        "id": log.id,
        "timestamp": _ts(log.timestamp),
        "level": log.level,
        "category": log.category,
        "action": log.action,
        "username": log.username,
        "full_name": full_name,
        "user_id": log.user_id,
        "ip_address": log.ip_address,
        "status": log.status or "success",
        "target_type": log.target_type,
        "target_id": log.target_id,
        "message": log.message,
        "details": log.details or {},
    }



@router.get("/alerts")
def list_alert_logs(
    start: Optional[str] = Query(None, description="ISO start datetime"),
    end: Optional[str] = Query(None, description="ISO end datetime"),
    server_id: Optional[int] = Query(None),
    level: Optional[str] = Query(None),
    keyword: Optional[str] = Query(None),
    regex: bool = Query(False),
    page: int = Query(1, ge=1),
    page_size: int = Query(DEFAULT_PAGE_SIZE, ge=1, le=500),
    authorization: Optional[str] = Header(None),
    db: Session = Depends(get_db),
):
    user = _require_logs(authorization, db, "view")
    allowed = _log_scope(db, user, server_id)
    query = _alert_base_query(db, start, end, server_id, level,
                               None if user.get("is_admin") else allowed)
    query, rx_kw, rx_fields = _apply_keyword_filter(query, Alert, keyword, regex)

    if regex and rx_kw:
        rows = query.order_by(Alert.timestamp.desc(), Alert.id.desc()).all()
        rows = _regex_filter_rows(rows, rx_kw, rx_fields)
        total = len(rows)
        start_idx = (page - 1) * page_size
        page_rows = rows[start_idx:start_idx + page_size]
    else:
        total = query.count()
        page_rows = (
            query.order_by(Alert.timestamp.desc(), Alert.id.desc())
            .offset((page - 1) * page_size)
            .limit(page_size)
            .all()
        )

    return {
        "total": total,
        "page": page,
        "page_size": page_size,
        "items": [_serialize_alert(r) for r in page_rows],
    }


@router.get("/alerts/{log_id}/context")
def alert_log_context(
    log_id: int,
    start: Optional[str] = Query(None),
    end: Optional[str] = Query(None),
    server_id: Optional[int] = Query(None),
    level: Optional[str] = Query(None),
    keyword: Optional[str] = Query(None),
    regex: bool = Query(False),
    authorization: Optional[str] = Header(None),
    db: Session = Depends(get_db),
):
    user = _require_logs(authorization, db, "view")
    target = db.query(Alert).filter(Alert.id == log_id).first()
    if not target:
        raise HTTPException(status_code=404, detail="日志记录不存在")

    query = _alert_base_query(db, start, end, server_id, level)
    query, rx_kw, rx_fields = _apply_keyword_filter(query, Alert, keyword, regex)

    if regex and rx_kw:
        rows = query.order_by(Alert.timestamp.asc(), Alert.id.asc()).all()
        rows = _regex_filter_rows(rows, rx_kw, rx_fields)
    else:
        rows = query.order_by(Alert.timestamp.asc(), Alert.id.asc()).all()

    index = next((i for i, r in enumerate(rows) if r[0].id == target.id), -1)
    if index < 0:
        return {
            "before": [],
            "record": _serialize_alert((target, target.server.name if target.server else "")),
            "after": [],
        }

    before = rows[max(0, index - CONTEXT_WINDOW):index]
    after = rows[index + 1:index + 1 + CONTEXT_WINDOW]
    return {
        "before": [_serialize_alert(r) for r in before],
        "record": _serialize_alert(rows[index]),
        "after": [_serialize_alert(r) for r in after],
    }


@router.get("/alerts/export")
def export_alert_logs(
    start: Optional[str] = Query(None),
    end: Optional[str] = Query(None),
    server_id: Optional[int] = Query(None),
    level: Optional[str] = Query(None),
    keyword: Optional[str] = Query(None),
    regex: bool = Query(False),
    authorization: Optional[str] = Header(None),
    db: Session = Depends(get_db),
):
    user = _require_logs(authorization, db, "export")
    allowed = _log_scope(db, user, server_id)
    query = _alert_base_query(db, start, end, server_id, level,
                               None if user.get("is_admin") else allowed)
    query, rx_kw, rx_fields = _apply_keyword_filter(query, Alert, keyword, regex)

    if regex and rx_kw:
        rows = query.order_by(Alert.timestamp.asc(), Alert.id.asc()).all()
        rows = _regex_filter_rows(rows, rx_kw, rx_fields)
    else:
        rows = query.order_by(Alert.timestamp.asc(), Alert.id.asc()).all()

    wb = _build_excel_workbook(
        headers=["时间", "主机", "级别", "类型", "标题", "消息", "确认状态"],
        rows=[
            [
                _ts(r[0].timestamp),
                r[1] or "",
                r[0].level,
                r[0].metric_type,
                r[0].title,
                r[0].message,
                "已确认" if r[0].acknowledged else "未确认",
            ]
            for r in rows
        ],
        sheet_title="告警日志",
    )
    return _xlsx_response(wb, "alert_logs.xlsx")



@router.get("/operations")
def list_operation_logs(
    start: Optional[str] = Query(None),
    end: Optional[str] = Query(None),
    category: Optional[str] = Query(None),
    username: Optional[str] = Query(None),
    server_id: Optional[int] = Query(None, description="只返回与该主机关联的操作记录"),
    keyword: Optional[str] = Query(None),
    regex: bool = Query(False),
    page: int = Query(1, ge=1),
    page_size: int = Query(DEFAULT_PAGE_SIZE, ge=1, le=500),
    authorization: Optional[str] = Header(None),
    db: Session = Depends(get_db),
):
    user = _require_logs(authorization, db, "view")
    query = _op_base_query(db, start, end, category, username, server_id)
    query, rx_kw, rx_fields = _apply_keyword_filter(query, OperationLog, keyword, regex)

    if regex and rx_kw:
        rows = query.order_by(OperationLog.timestamp.desc(), OperationLog.id.desc()).all()
        rows = _regex_filter_rows(rows, rx_kw, rx_fields)
        total = len(rows)
        start_idx = (page - 1) * page_size
        page_rows = rows[start_idx:start_idx + page_size]
    else:
        total = query.count()
        page_rows = (
            query.order_by(OperationLog.timestamp.desc(), OperationLog.id.desc())
            .offset((page - 1) * page_size)
            .limit(page_size)
            .all()
        )

    user_map = _build_user_map(db, page_rows)
    return {
        "total": total,
        "page": page,
        "page_size": page_size,
        "items": [_serialize_op(r, user_map) for r in page_rows],
    }


@router.get("/operations/{log_id}/context")
def operation_log_context(
    log_id: int,
    start: Optional[str] = Query(None),
    end: Optional[str] = Query(None),
    category: Optional[str] = Query(None),
    username: Optional[str] = Query(None),
    server_id: Optional[int] = Query(None),
    keyword: Optional[str] = Query(None),
    regex: bool = Query(False),
    authorization: Optional[str] = Header(None),
    db: Session = Depends(get_db),
):
    user = _require_logs(authorization, db, "view")
    target = db.query(OperationLog).filter(OperationLog.id == log_id).first()
    if not target:
        raise HTTPException(status_code=404, detail="日志记录不存在")

    query = _op_base_query(db, start, end, category, username, server_id)
    query, rx_kw, rx_fields = _apply_keyword_filter(query, OperationLog, keyword, regex)

    if regex and rx_kw:
        rows = query.order_by(OperationLog.timestamp.asc(), OperationLog.id.asc()).all()
        rows = _regex_filter_rows(rows, rx_kw, rx_fields)
    else:
        rows = query.order_by(OperationLog.timestamp.asc(), OperationLog.id.asc()).all()

    index = next((i for i, r in enumerate(rows) if r.id == target.id), -1)
    if index < 0:
        return {
            "before": [],
            "record": _serialize_op(target),
            "after": [],
        }

    before = rows[max(0, index - CONTEXT_WINDOW):index]
    after = rows[index + 1:index + 1 + CONTEXT_WINDOW]
    user_map = _build_user_map(db, rows)
    return {
        "before": [_serialize_op(r, user_map) for r in before],
        "record": _serialize_op(rows[index], user_map),
        "after": [_serialize_op(r, user_map) for r in after],
    }


@router.get("/operations/export")
def export_operation_logs(
    start: Optional[str] = Query(None),
    end: Optional[str] = Query(None),
    category: Optional[str] = Query(None),
    username: Optional[str] = Query(None),
    server_id: Optional[int] = Query(None),
    keyword: Optional[str] = Query(None),
    regex: bool = Query(False),
    authorization: Optional[str] = Header(None),
    db: Session = Depends(get_db),
):
    user = _require_logs(authorization, db, "export")
    query = _op_base_query(db, start, end, category, username, server_id)
    query, rx_kw, rx_fields = _apply_keyword_filter(query, OperationLog, keyword, regex)

    if regex and rx_kw:
        rows = query.order_by(OperationLog.timestamp.asc(), OperationLog.id.asc()).all()
        rows = _regex_filter_rows(rows, rx_kw, rx_fields)
    else:
        rows = query.order_by(OperationLog.timestamp.asc(), OperationLog.id.asc()).all()

    user_map = _build_user_map(db, rows)
    wb = _build_excel_workbook(
        headers=["时间", "级别", "分类", "动作", "状态", "用户", "来源IP", "目标类型", "目标ID", "消息", "详情"],
        rows=[
            [
                _ts(r.timestamp),
                r.level,
                r.category,
                r.action,
                "成功" if (r.status or "success") == "success" else "失败",
                user_map.get(r.username) or r.username or "—",
                r.ip_address or "—",
                r.target_type,
                r.target_id,
                r.message,
                _details_str(r.details),
            ]
            for r in rows
        ],
        sheet_title="操作日志",
    )
    start_str = _filename_dt(start)
    end_str = _filename_dt(end)
    if start_str and end_str:
        filename = f"操作日志 {start_str}-{end_str}.xlsx"
    elif start_str:
        filename = f"操作日志 {start_str}起.xlsx"
    elif end_str:
        filename = f"操作日志 截至{end_str}.xlsx"
    else:
        filename = "操作日志.xlsx"
    return _xlsx_response(wb, filename)



def _details_str(details):
    if not details:
        return ""
    try:
        import json
        return json.dumps(details, ensure_ascii=False)
    except Exception:
        return str(details)


def _build_excel_workbook(headers, rows, sheet_title):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="0275D8")
    thin = Side(style="thin", color="D0D7DE")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.append(headers)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    for row in rows:
        ws.append(row)

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                max_len = max(max_len, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 2, 60)

    return wb


def _xlsx_response(wb, filename):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    ascii_name = filename.encode('ascii', 'ignore').decode() or "logs.xlsx"
    encoded_name = quote(filename, safe='')
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": (
                f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{encoded_name}"
            )
        },
    )




class ClearLogRequest(BaseModel):
    start: Optional[str] = None
    end: Optional[str] = None
    server_id: Optional[int] = None
    level: Optional[str] = None
    category: Optional[str] = None
    username: Optional[str] = None
    keyword: Optional[str] = None
    regex: bool = False


def _protected_cutoff() -> datetime:
    return datetime.now(timezone.utc) - timedelta(hours=ACTIVE_LOG_PROTECTION_HOURS)


def _apply_filters_to_query(query, model, start, end, keyword, regex):
    query, rx_kw, rx_fields = _apply_keyword_filter(query, model, keyword, regex)
    if regex and rx_kw:
        return query, rx_kw, rx_fields
    return query, None, None


@router.delete("/alerts/clear")
def clear_alert_logs(
    request: Request,
    data: ClearLogRequest,
    authorization: Optional[str] = Header(None),
    db: Session = Depends(get_db),
):
    user = require_admin_full(authorization, db)
    protected = _protected_cutoff()

    query = _alert_base_query(db, data.start, data.end, data.server_id, data.level)
    query, rx_kw, rx_fields = _apply_filters_to_query(query, Alert, data.start, data.end, data.keyword, data.regex)

    # Apply active-log protection: never delete logs newer than the protection window.
    query = query.filter(Alert.timestamp < protected)

    if rx_kw and rx_fields:
        rows = query.order_by(Alert.timestamp.asc(), Alert.id.asc()).all()
        rows = _regex_filter_rows(rows, rx_kw, rx_fields)
        ids = [r[0].id for r in rows]
        deleted = 0
        if ids:
            batch = ids[:CLEAR_MAX_LIMIT]
            deleted = db.query(Alert).filter(Alert.id.in_(batch)).delete(synchronize_session=False)
    else:
        count = min(query.count(), CLEAR_MAX_LIMIT)
        ids = (
            query.order_by(Alert.timestamp.asc(), Alert.id.asc())
            .limit(count)
            .with_entities(Alert.id)
            .all()
        )
        deleted = 0
        if ids:
            deleted = db.query(Alert).filter(Alert.id.in_([i[0] for i in ids])).delete(synchronize_session=False)

    db.commit()

    log_operation(
        db,
        category="log",
        action="clear",
        username=user.get("username", ""),
        user_id=user.get("user_id"),
        ip_address=_client_ip(request),
        target_type="alerts",
        target_id="",
        message=f"用户 {user.get('username', '')} 清除告警日志 {deleted} 条",
        details={
            "type": "alerts",
            "deleted": deleted,
            "filters": data.model_dump(exclude_none=True),
            "protected_until": protected.isoformat(),
        },
    )

    return {
        "deleted": deleted,
        "protected_until": protected.isoformat(),
        "message": f"已清除 {deleted} 条告警日志（最近 {ACTIVE_LOG_PROTECTION_HOURS} 小时内的活跃日志已保留）",
    }


@router.delete("/operations/clear")
def clear_operation_logs(
    request: Request,
    data: ClearLogRequest,
    authorization: Optional[str] = Header(None),
    db: Session = Depends(get_db),
):
    user = require_admin_full(authorization, db)
    protected = _protected_cutoff()

    query = _op_base_query(db, data.start, data.end, data.category, data.username)
    query, rx_kw, rx_fields = _apply_filters_to_query(query, OperationLog, data.start, data.end, data.keyword, data.regex)
    query = query.filter(OperationLog.timestamp < protected)

    if rx_kw and rx_fields:
        rows = query.order_by(OperationLog.timestamp.asc(), OperationLog.id.asc()).all()
        rows = _regex_filter_rows(rows, rx_kw, rx_fields)
        ids = [r.id for r in rows]
        deleted = 0
        if ids:
            batch = ids[:CLEAR_MAX_LIMIT]
            deleted = db.query(OperationLog).filter(OperationLog.id.in_(batch)).delete(synchronize_session=False)
    else:
        count = min(query.count(), CLEAR_MAX_LIMIT)
        ids = (
            query.order_by(OperationLog.timestamp.asc(), OperationLog.id.asc())
            .limit(count)
            .with_entities(OperationLog.id)
            .all()
        )
        deleted = 0
        if ids:
            deleted = db.query(OperationLog).filter(OperationLog.id.in_([i[0] for i in ids])).delete(synchronize_session=False)

    db.commit()

    log_operation(
        db,
        category="log",
        action="clear",
        username=user.get("username", ""),
        user_id=user.get("user_id"),
        ip_address=_client_ip(request),
        target_type="operations",
        target_id="",
        message=f"用户 {user.get('username', '')} 清除操作日志 {deleted} 条",
        details={
            "type": "operations",
            "deleted": deleted,
            "filters": data.model_dump(exclude_none=True),
            "protected_until": protected.isoformat(),
        },
    )

    return {
        "deleted": deleted,
        "protected_until": protected.isoformat(),
        "message": f"已清除 {deleted} 条操作日志（最近 {ACTIVE_LOG_PROTECTION_HOURS} 小时内的活跃日志已保留）",
    }
